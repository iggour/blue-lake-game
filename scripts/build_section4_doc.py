from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path('/Users/zhangtongyang/Desktop/SSAT/section4.docx')

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_text(cell, text, bold=False, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'STHeiti'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'STHeiti')
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def math_run(parent, text):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = text
    r.append(t)
    parent.append(r)

def math_text(parent, text):
    math_run(parent, text)

def math_fraction(parent, numerator, denominator):
    f = OxmlElement('m:f')
    num = OxmlElement('m:num')
    den = OxmlElement('m:den')
    math_text(num, numerator)
    math_text(den, denominator)
    f.append(num)
    f.append(den)
    parent.append(f)

def math_sup(parent, base, exponent):
    s = OxmlElement('m:sSup')
    e = OxmlElement('m:e')
    sup = OxmlElement('m:sup')
    math_text(e, base)
    math_text(sup, exponent)
    s.append(e)
    s.append(sup)
    parent.append(s)

def add_math(p, parts):
    omath = OxmlElement('m:oMath')
    for part in parts:
        if isinstance(part, tuple) and part[0] == 'frac':
            math_fraction(omath, part[1], part[2])
        elif isinstance(part, tuple) and part[0] == 'sup':
            math_sup(omath, part[1], part[2])
        else:
            math_text(omath, str(part))
    p._p.append(omath)

def add_formula(doc, label, parts):
    p = doc.add_paragraph(style='Formula')
    if label:
        r = p.add_run(label + '  ')
        r.bold = True
    add_math(p, parts)
    return p

def add_bullet(doc, text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def add_topic(doc, title, tag):
    p = doc.add_paragraph(style='Topic')
    r = p.add_run(title)
    r.bold = True
    p.add_run('  ' + tag)
    return p

def add_section_heading(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.add_run(text)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(4)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1']._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
styles['Heading 1'].font.size = Pt(15)
styles['Heading 1'].font.color.rgb = RGBColor(31, 78, 121)
styles['Heading 2'].font.name = 'Aptos'
styles['Heading 2']._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].font.color.rgb = RGBColor(31, 78, 121)

if 'Formula' not in styles:
    styles.add_style('Formula', 1)
styles['Formula'].font.name = 'Cambria Math'
styles['Formula']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
styles['Formula'].font.size = Pt(12)
styles['Formula'].paragraph_format.left_indent = Inches(0.25)
styles['Formula'].paragraph_format.space_after = Pt(6)

if 'Topic' not in styles:
    styles.add_style('Topic', 1)
styles['Topic'].font.name = 'Aptos'
styles['Topic']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Topic']._element.rPr.rFonts.set(qn('w:eastAsia'), 'STHeiti')
styles['Topic'].font.size = Pt(10.5)
styles['Topic'].paragraph_format.space_before = Pt(4)
styles['Topic'].paragraph_format.space_after = Pt(2)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('Section 4')
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(31, 78, 121)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run('核心公式与算法速查  |  Knowledge Base')
r.bold = True
r.font.size = Pt(13)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run('SSAT 数学常用公式、解题算法与术语对照').italic = True

add_section_heading(doc, '第一部分：核心公式与算法速查')

doc.add_paragraph('1. Numbers and Operations（数与运算）', style='Heading 2')
add_topic(doc, 'Fraction Addition and Subtraction（分数加减法）', 'Q3, Q24')
add_bullet(doc, 'Algorithm: Find a common denominator, then add or subtract the numerators.')
add_formula(doc, 'Formula:', [('frac','a','b'), ' ± ', ('frac','c','d'), ' = ', ('frac','ad ± bc','bd')])
add_formula(doc, 'Example:', [('frac','3','6'), ' − ', ('frac','1','4'), ' = ', ('frac','2','4'), ' − ', ('frac','1','4'), ' = ', ('frac','1','4')])

add_topic(doc, 'Fraction Multiplication and Division（分数乘除法）', 'Q24')
add_bullet(doc, 'Multiplication: Multiply numerators together and denominators together.')
add_formula(doc, 'Multiplication:', [('frac','a','b'), ' × ', ('frac','c','d'), ' = ', ('frac','ac','bd')])
add_bullet(doc, 'Division: Multiply by the reciprocal of the divisor.')
add_formula(doc, 'Division:', [('frac','a','b'), ' ÷ ', ('frac','c','d'), ' = ', ('frac','a','b'), ' × ', ('frac','d','c')])

add_topic(doc, 'Percentage Calculation（百分比计算）', 'Q9, Q19')
add_formula(doc, 'Formula:', ['Part = Whole × Percent'])
add_formula(doc, 'Conversion:', ['p% = ', ('frac','p','100')])

add_topic(doc, 'Estimation（估算技巧）', 'Q11')
add_bullet(doc, 'Algorithm: Round complex numbers to the nearest ten, hundred, or thousand to quickly eliminate incorrect options.')

doc.add_paragraph('2. Algebra（代数初步）', style='Heading 2')
add_topic(doc, 'Algebraic Expressions（代数式表达）', 'Q7, Q17')
add_formula(doc, 'Translation:', ['“A is 4 more than B”  ⇒  A = B + 4'])
add_formula(doc, 'Tiered Pricing:', ['Total Cost = Base Cost + (Total Quantity − Base Quantity) × Additional Rate'])

add_topic(doc, 'Solving Linear Equations（解一元一次方程）', 'Q20')
add_bullet(doc, 'Algorithm: Isolate the variable on one side and constants on the other.')
add_formula(doc, 'Example:', ['3 + 7n = 6  ⇒  7n = 6 − 3  ⇒  7n = 3  ⇒  n = ', ('frac','3','7')])

add_topic(doc, 'Number Sequences（数列规律）', 'Q5')
add_bullet(doc, 'Algorithm: Observe the differences or ratios between adjacent terms, or the independent patterns of the numerator and denominator.')

doc.add_paragraph('3. Geometry（几何与空间）', style='Heading 2')
add_topic(doc, 'Rectangle and Square（矩形与正方形）', 'Q16')
add_formula(doc, 'Square Area:', ['A = ', ('sup','s','2')])
add_formula(doc, 'Rectangle Area:', ['A = l × w'])
add_formula(doc, 'Perimeter:', ['P = 2(l + w)'])

add_topic(doc, 'Angles with Parallel Lines（平行线角度）', 'Q23')
add_bullet(doc, 'When lines are parallel: corresponding angles are equal, alternate interior angles are equal, and consecutive interior angles sum to 180°.')

add_topic(doc, '3D Views（立体图形三视图）', 'Q12')
add_bullet(doc, "Algorithm: Based on the top view and the sphere's diameter, determine the length, width, and height of the rectangular package.")
add_formula(doc, 'Diameter:', ['d = 2r'])

doc.add_paragraph('4. Statistics and Word Problems（统计与应用题）', style='Heading 2')
add_topic(doc, 'Median（中位数）', 'Q15')
add_bullet(doc, 'Algorithm: Arrange all numbers in ascending order. If the count is odd, take the middle number. If the count is even, take the average of the two middle numbers.')

add_topic(doc, 'Rate and Time（速率问题）', 'Q14')
add_formula(doc, 'Formula:', ['Total Work = Rate × Time'])
add_formula(doc, 'Conversion:', ['1 hour = 60 minutes'])

add_topic(doc, 'Decimal Division Application（小数除法应用）', 'Q22')
add_bullet(doc, 'Algorithm: Calculate how many complete items can be made by taking the integer part, then calculate the remainder.')

add_section_heading(doc, '附录：本试卷常用数学术语中英对照')
terms = [
    ('Integer', '整数'), ('Prime number', '质数'), ('Median', '中位数'),
    ('Numerator', '分子'), ('Denominator', '分母'), ('Rectangle', '长方形'),
    ('Square', '正方形'), ('Radius', '半径'), ('Diameter', '直径'),
    ('Intersection', '交叉路口 / 交点'), ('Parallel', '平行的'),
    ('Fraction', '分数'), ('Percentage', '百分比'), ('Equation', '方程'),
    ('Sequence', '数列'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
set_cell_shading(table.rows[0].cells[0], '1F4E79')
set_cell_shading(table.rows[0].cells[1], '1F4E79')
set_cell_text(table.rows[0].cells[0], 'English', True, 'FFFFFF')
set_cell_text(table.rows[0].cells[1], '中文', True, 'FFFFFF')
for en, zh in terms:
    cells = table.add_row().cells
    set_cell_text(cells[0], en)
    set_cell_text(cells[1], zh)

for p in doc.paragraphs:
    p.paragraph_format.keep_together = True

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
